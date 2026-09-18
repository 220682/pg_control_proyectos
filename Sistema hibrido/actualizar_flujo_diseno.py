"""Reescribe Flujo y diseño.xlsx para alinearlo con control_de_proyectos.txt v5
y el spec workspace-nav-proyecto (2026-08-17). Elimina columnas duplicadas
y corrige typos / numeración confusa."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent

# Catálogo oficial (sección 5.1 control_de_proyectos.txt v5)
CATALOGO = [
    ("001", "Presupuesto del proyecto", "Al inicio", "Jefe de Of. Tecnica", "Existe (subida manual)"),
    ("002", "DP (Datos del proyecto)", "Al inicio", "Sup. Of. Tecnica", "Construido (import Excel)"),
    ("003", "PR (Reporte del proyecto)", "Al inicio", "Sup. Of. Tecnica", "Construido (auto con DP)"),
    ("004", "Listado recursos hh, hm, mat. (sin costo)", "Al inicio", "Sup. Of. Tecnica", "No creado"),
    ("005", "Listado materiales con costo", "Al inicio", "Sup. Costos", "No creado"),
    ("006", "Listado estimado de PETS", "Al inicio", "Sup. Of. Tecnica", "No creado"),
    ("007", "Alcance del proyecto", "Al inicio", "Sup. Costos", "No creado"),
    ("008", "Requerimiento del proyecto", "Al inicio y durante", "Sup.Op + Of.Tec.", "No creado"),
    ("009", "Listado de Personal requerido", "Al inicio", "Sup. Of. Tecnica", "No creado"),
    ("010", "RDT (Reporte Diario de Trabajo)", "Durante", "Sup. Operativo", "Existe (Excel; sin web)"),
    ("011", "Sistema LPS / Cronograma", "Durante", "Planner", "En uso (Excel multi-proyecto)"),
    ("012", "Programacion diaria", "Durante", "Planner", "Nuevo (multi-proyecto)"),
    ("013", "Status de requerimiento", "Durante", "Sup. Logistica", "Nuevo"),
    ("014", "Registro de gastos por RQ", "Durante", "Sup. Logistica", "Nuevo"),
    ("015", "Programacion de capacitaciones", "Durante", "Planner", "Nuevo"),
    ("016", "Status de capacitacion", "Durante", "Sup. Admin.", "Nuevo"),
    ("017", "Tareo diario personal indirecto", "Durante", "Sup. Admin.", "Nuevo"),
    ("018", "Status de tareo MOI", "Durante", "Sup. Admin.", "Nuevo"),
    ("019", "Informe de proyecto", "Cierre", "Sup. Operativo", "Nuevo"),
    ("020", "Acta de conformidad", "Cierre", "Sup. Operativo", "Nuevo"),
    ("021", "Status de Proyectos (dashboard)", "Transversal", "Sistema (PR)", "Pendiente"),
    ("--", "Costo de RQ de proyectos (semanal)", "Durante", "Sup. Logistica", "Pendiente"),
    ("022", "Orden de Trabajo (OT)", "Al inicio", "Jefe de Of. Tecnica", "Spec; sin construir"),
]

# Navegacion web dentro de un proyecto (spec 2026-08-17, nav 001-024)
NAV_GRUPOS: list[tuple[str, list[tuple[str, str, str]]]] = [
    (
        "Proyecto (sidebar izquierdo)",
        [
            ("001", "Dashboard", "Si — /proyectos/{id}"),
            ("002", "(OT) Orden de trabajo", "—"),
            ("003", "Presupuesto del proyecto", "—"),
            ("004", "DP (Datos del proyecto)", "Si — /proyectos/{id}/dp"),
            ("005", "Recursos hh, hm, mat. (s/c)", "—"),
            ("006", "Materiales (c/c)", "—"),
            ("007", "Alcance del servicio", "— (= cat. 007 Alcance del proyecto)"),
            ("008", "Lista de personal nuevo", "— (= cat. 009 Personal requerido)"),
            ("009", "Consolidado de servicio", "— (futuro; relacion OT)"),
        ],
    ),
    (
        "Planificacion (panel derecho)",
        [
            ("010", "3WLA / Sistema LPS", "— (= cat. 011 Cronograma LPS)"),
            ("011", "Programacion diaria", "— (= cat. 012)"),
            ("012", "PR (Reporte del proyecto)", "Si — /proyectos/{id}/pr"),
            ("013", "Status de capacitaciones", "— (= cat. 016)"),
            ("014", "Programacion de capacitaciones", "— (= cat. 015)"),
        ],
    ),
    (
        "Costos",
        [("015", "Status de servicios", "—")],
    ),
    (
        "Supervision operativa",
        [
            ("016", "RDT (Registro diario de trabajo)", "— (= cat. 010)"),
            ("017", "Informe de servicio", "— (= cat. 019 Informe de proyecto)"),
            ("018", "Requerimiento del servicio", "— (= cat. 008)"),
            ("019", "Acta de conformidad", "— (= cat. 020)"),
        ],
    ),
    (
        "Logistica",
        [
            ("020", "Status de requerimiento", "— (= cat. 013)"),
            ("021", "Registro de costos por servicios", "— (= cat. 014 + Costo RQ)"),
        ],
    ),
    (
        "Administracion",
        [
            ("022", "Tareo MOI", "— (= cat. 017)"),
            ("023", "Consolidado MOI", "— (= cat. 018)"),
        ],
    ),
    (
        "SSOMA",
        [
            (
                "024",
                "Pets (repositorio unico contratista)",
                "Portafolio — NO es listado por proyecto",
            ),
        ],
    ),
]

PORTAFOLIO = [
    ("021", "Status de Proyectos (dashboard)", "Vista portafolio — multi-proyecto"),
    ("—", "+ Nuevo Proyecto", "Primera fila del desplegable de proyectos"),
    ("024", "Pets / Cargar PETS", "Repositorio unico SSOMA"),
]

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(bold=True, color="E2E8F0", size=11)
TITLE_FONT = Font(bold=True, size=12)
NOTE_FONT = Font(italic=True, color="64748B", size=10)
GROUP_FILL = PatternFill("solid", fgColor="334155")
GROUP_FONT = Font(bold=True, color="94A3B8", size=10)


def encontrar_excel() -> Path:
    for f in BASE.iterdir():
        if f.suffix == ".xlsx" and f.name.startswith("Flujo") and not f.name.startswith("~$"):
            return f
    raise FileNotFoundError("No se encontro Flujo y diseno.xlsx")


def escribir_fila(ws, row: int, valores: list, bold: bool = False, col_inicio: int = 1) -> None:
    for offset, val in enumerate(valores):
        cell = ws.cell(row=row, column=col_inicio + offset, value=val)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            cell.font = Font(bold=True)


def main() -> None:
    path = encontrar_excel()
    lock = path.parent / f"~${path.name}"
    if lock.exists():
        raise SystemExit(f"CERRAR EXCEL PRIMERO: existe {lock.name}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Referencia"

    # --- Titulo y nota explicativa ---
    ws.merge_cells("A1:J1")
    ws["A1"] = "Flujo y diseno — Referencia unica (v2, 2026-08-19)"
    ws["A1"].font = TITLE_FONT

    ws.merge_cells("A2:J2")
    ws["A2"] = (
        "Dos numeraciones distintas a proposito: (A) CATALOGO 001-022 = documentos del sistema "
        "(control_de_proyectos.txt). (B) NAV 001-024 = menu de la web (spec workspace-nav-proyecto). "
        "No mezclar columnas — la version anterior tenia 3 copias del mismo menu y numeros cruzados."
    )
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = Alignment(wrap_text=True)

    row = 4
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row=row, column=1, value="A. CATALOGO DE FORMATOS (documentos del sistema)").font = TITLE_FONT
    row += 1
    escribir_fila(ws, row, ["N. cat.", "Formato", "Fase", "Responsable", "Estado web"], bold=True)
    for r in range(1, 6):
        ws.cell(row=row, column=r).fill = HEADER_FILL
        ws.cell(row=row, column=r).font = HEADER_FONT
    row += 1
    for item in CATALOGO:
        escribir_fila(ws, row, list(item))
        row += 1

    row += 1
    nav_col = 7
    ws.merge_cells(start_row=row, start_column=nav_col, end_row=row, end_column=nav_col + 2)
    ws.cell(row=row, column=nav_col, value="B. NAVEGACION WEB — dentro de un proyecto").font = TITLE_FONT
    nav_header_row = row + 1
    escribir_fila(ws, nav_header_row, ["N. nav", "Item", "Pantalla real"], bold=True, col_inicio=nav_col)
    for c in range(nav_col, nav_col + 3):
        cell = ws.cell(row=nav_header_row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    nav_row = nav_header_row + 1
    for titulo_grupo, items in NAV_GRUPOS:
        ws.merge_cells(start_row=nav_row, start_column=nav_col, end_row=nav_row, end_column=nav_col + 2)
        gcell = ws.cell(row=nav_row, column=nav_col, value=titulo_grupo)
        gcell.fill = GROUP_FILL
        gcell.font = GROUP_FONT
        nav_row += 1
        for num, nombre, estado in items:
            escribir_fila(ws, nav_row, [num, nombre, estado], col_inicio=nav_col)
            nav_row += 1

    row = nav_row + 2
    ws.merge_cells(start_row=row, start_column=nav_col, end_row=row, end_column=nav_col + 2)
    ws.cell(row=row, column=nav_col, value="C. PORTAFOLIO (fuera de un proyecto)").font = TITLE_FONT
    row += 1
    escribir_fila(ws, row, ["Ref.", "Item", "Nota"], bold=True, col_inicio=nav_col)
    for c in range(nav_col, nav_col + 3):
        ws.cell(row=row, column=c).fill = HEADER_FILL
        ws.cell(row=row, column=c).font = HEADER_FONT
    row += 1
    for item in PORTAFOLIO:
        escribir_fila(ws, row, list(item), col_inicio=nav_col)
        row += 1

    # Anchos
    for col, width in enumerate([8, 42, 18, 22, 28, 3, 8, 38, 36, 5], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(path)
    print(f"OK: {path}")


if __name__ == "__main__":
    main()
